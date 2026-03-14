from fastapi import FastAPI, APIRouter, Query, HTTPException
from fastapi.responses import StreamingResponse
from dotenv import load_dotenv
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
import os
import logging
import asyncio
import random
import io
import csv as csv_module
import string
from pathlib import Path
from pydantic import BaseModel, Field, ConfigDict
from typing import List, Optional
import uuid
from datetime import datetime, timezone, timedelta

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

app = FastAPI()
api_router = APIRouter(prefix="/api")

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(name)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)

# ═══════════════════════════════════════════════════
# MODELS
# ═══════════════════════════════════════════════════

class WorkerStartRequest(BaseModel):
    marketplace: str = "amazon.de"
    mode: str = "continuous"
    target_count: int = 0
    min_delay: int = 25
    max_delay: int = 45
    manual_review: bool = False

class SettingsUpdate(BaseModel):
    google_sheets_enabled: Optional[bool] = None
    google_sheets_url: Optional[str] = None
    manual_review_enabled: Optional[bool] = None
    default_marketplace: Optional[str] = None
    default_min_delay: Optional[int] = None
    default_max_delay: Optional[int] = None
    default_mode: Optional[str] = None
    default_target_count: Optional[int] = None
    scoring_weights: Optional[dict] = None
    duplicate_rules: Optional[dict] = None
    export_options: Optional[dict] = None

class NotesUpdate(BaseModel):
    notes: str

class FlagUpdate(BaseModel):
    flagged: bool = True
    flagged_reason: str = ""

class SavedViewCreate(BaseModel):
    name: str
    filters: dict = {}

EU_COUNTRIES = {"Germany", "Spain", "France", "Italy", "Netherlands", "Belgium", "Austria",
    "Portugal", "Greece", "Poland", "Sweden", "Denmark", "Finland", "Ireland",
    "Czech Republic", "Romania", "Hungary", "Croatia", "Slovakia", "Slovenia",
    "Bulgaria", "Lithuania", "Latvia", "Estonia", "Luxembourg", "Malta", "Cyprus"}

DEFAULT_SCORING_WEIGHTS = {
    "vat_present": 20, "registration_present": 15, "phone_present": 20,
    "email_present": 20, "address_present": 15, "eu_country": 10,
}

def calculate_lead_score(seller, weights=None):
    w = weights or DEFAULT_SCORING_WEIGHTS
    score = 0
    if seller.get("vat_number_if_visible"):
        score += w.get("vat_present", 20)
    if seller.get("registration_number_if_visible"):
        score += w.get("registration_present", 15)
    if seller.get("public_phone_if_visible"):
        score += w.get("phone_present", 20)
    if seller.get("public_email_if_visible"):
        score += w.get("email_present", 20)
    if seller.get("business_address_if_visible"):
        score += w.get("address_present", 15)
    if seller.get("country_if_visible") in EU_COUNTRIES:
        score += w.get("eu_country", 10)
    return min(score, 100)

def get_score_breakdown(seller, weights=None):
    w = weights or DEFAULT_SCORING_WEIGHTS
    return [
        {"factor": "VAT Number", "present": bool(seller.get("vat_number_if_visible")), "points": w.get("vat_present", 20)},
        {"factor": "Registration #", "present": bool(seller.get("registration_number_if_visible")), "points": w.get("registration_present", 15)},
        {"factor": "Phone", "present": bool(seller.get("public_phone_if_visible")), "points": w.get("phone_present", 20)},
        {"factor": "Email", "present": bool(seller.get("public_email_if_visible")), "points": w.get("email_present", 20)},
        {"factor": "Address", "present": bool(seller.get("business_address_if_visible")), "points": w.get("address_present", 15)},
        {"factor": "EU Country", "present": seller.get("country_if_visible") in EU_COUNTRIES, "points": w.get("eu_country", 10)},
    ]

# ═══════════════════════════════════════════════════
# MOCK DATA CONSTANTS
# ═══════════════════════════════════════════════════

PRODUCT_TITLES = [
    "Wireless Bluetooth Headphones Noise Cancelling Over Ear",
    "USB C Charging Cable Fast Charge 3-Pack 2m",
    "Portable Power Bank 20000mAh Quick Charge USB",
    "Smart Watch Fitness Tracker Heart Rate Monitor",
    "LED Desk Lamp Dimmable Touch Control USB Port",
    "Mechanical Keyboard RGB Backlit Gaming Compact",
    "Webcam 1080p HD Streaming USB Microphone Built-in",
    "Wireless Mouse Ergonomic Silent Click Rechargeable",
    "Phone Stand Adjustable Aluminum Desktop Holder Universal",
    "Laptop Sleeve Case 15.6 Inch Waterproof Shockproof",
    "Bluetooth Speaker Portable Waterproof 24H Playtime",
    "HDMI Cable 4K 60Hz High Speed Braided 2m",
    "Ring Light 10 Inch LED Dimmable Tripod Stand",
    "Air Purifier HEPA Filter Quiet Sleep Mode",
    "Electric Kettle Temperature Control Stainless Steel 1.7L",
    "Yoga Mat Non Slip Eco Friendly 6mm Thick",
    "Resistance Bands Set Exercise Workout 5-Pack Latex",
    "Kitchen Scale Digital Precision Stainless Steel",
    "Water Bottle Stainless Steel Insulated 750ml BPA-Free",
    "Car Phone Mount Magnetic Air Vent Universal Holder",
]

MARKETPLACE_CONFIG = {
    "amazon.de": {
        "country": "Germany",
        "domain": "amazon.de",
        "vat_prefix": "DE",
        "phone_prefix": "+49",
        "email_tld": "de",
        "sellers": [
            {"name": "TechGear Berlin", "biz": "TechGear GmbH", "city": "Berlin", "street": "Friedrichstr. 123, 10117"},
            {"name": "ElektroHaus Muenchen", "biz": "ElektroHaus AG", "city": "Muenchen", "street": "Leopoldstr. 45, 80802"},
            {"name": "BueroProfi Hamburg", "biz": "BueroProfi e.K.", "city": "Hamburg", "street": "Moenckebergstr. 7, 20095"},
            {"name": "DigitalWelt Frankfurt", "biz": "DigitalWelt GmbH", "city": "Frankfurt", "street": "Zeil 106, 60313"},
            {"name": "SmartHome Koeln", "biz": "SmartHome Solutions GmbH", "city": "Koeln", "street": "Hohe Str. 52, 50667"},
            {"name": "GreenTech Stuttgart", "biz": "GreenTech Handel GmbH", "city": "Stuttgart", "street": "Koenigstr. 28, 70173"},
            {"name": "ProAudio Duesseldorf", "biz": "ProAudio Deutschland GmbH", "city": "Duesseldorf", "street": "Schadowstr. 89, 40212"},
            {"name": "FitnessPro Leipzig", "biz": "FitnessPro GmbH", "city": "Leipzig", "street": "Petersstr. 15, 04109"},
            {"name": "KuechenKoenig Dresden", "biz": "KuechenKoenig e.K.", "city": "Dresden", "street": "Prager Str. 2, 01069"},
            {"name": "AutoZubehoer Nuernberg", "biz": "AutoZubehoer Bayern GmbH", "city": "Nuernberg", "street": "Karolinenstr. 34, 90402"},
            {"name": "SportDirekt Bremen", "biz": "SportDirekt Handels GmbH", "city": "Bremen", "street": "Obernstr. 18, 28195"},
            {"name": "LichtDesign Hannover", "biz": "LichtDesign Nord GmbH", "city": "Hannover", "street": "Georgstr. 44, 30159"},
        ],
    },
    "amazon.es": {
        "country": "Spain",
        "domain": "amazon.es",
        "vat_prefix": "ES",
        "phone_prefix": "+34",
        "email_tld": "es",
        "sellers": [
            {"name": "TecnoTienda Madrid", "biz": "TecnoTienda S.L.", "city": "Madrid", "street": "Gran Via 28, 28013"},
            {"name": "ElectroShop Barcelona", "biz": "ElectroShop S.A.", "city": "Barcelona", "street": "Passeig de Gracia 55, 08007"},
            {"name": "InfoTech Valencia", "biz": "InfoTech Iberia S.L.", "city": "Valencia", "street": "Calle Colon 48, 46004"},
            {"name": "CasaSmart Sevilla", "biz": "CasaSmart Andalucia S.L.", "city": "Sevilla", "street": "Av. Constitucion 20, 41001"},
            {"name": "GadgetZone Malaga", "biz": "GadgetZone Costa S.L.", "city": "Malaga", "street": "Calle Larios 12, 29015"},
            {"name": "OfiPro Bilbao", "biz": "OfiPro Norte S.L.", "city": "Bilbao", "street": "Gran Via 8, 48001"},
            {"name": "AudioPremium Zaragoza", "biz": "AudioPremium Aragon S.L.", "city": "Zaragoza", "street": "Paseo Independencia 24, 50004"},
            {"name": "DeporteMax Alicante", "biz": "DeporteMax Levante S.L.", "city": "Alicante", "street": "Av. Maisonnave 33, 03003"},
            {"name": "CocinaPlus Murcia", "biz": "CocinaPlus Sur S.L.", "city": "Murcia", "street": "Gran Via Salzillo 15, 30004"},
            {"name": "LuzHogar Palma", "biz": "LuzHogar Baleares S.L.", "city": "Palma", "street": "Passeig del Born 7, 07012"},
        ],
    },
    "amazon.fr": {
        "country": "France",
        "domain": "amazon.fr",
        "vat_prefix": "FR",
        "phone_prefix": "+33",
        "email_tld": "fr",
        "sellers": [
            {"name": "TechBoutique Paris", "biz": "TechBoutique SARL", "city": "Paris", "street": "25 Rue de Rivoli, 75001"},
            {"name": "ElectroFrance Lyon", "biz": "ElectroFrance SAS", "city": "Lyon", "street": "18 Rue de la Republique, 69002"},
            {"name": "NumeriPro Marseille", "biz": "NumeriPro Mediterranee SARL", "city": "Marseille", "street": "65 La Canebiere, 13001"},
            {"name": "MaisonConnect Toulouse", "biz": "MaisonConnect Occitanie SAS", "city": "Toulouse", "street": "32 Rue Alsace-Lorraine, 31000"},
            {"name": "AudioPlus Bordeaux", "biz": "AudioPlus Aquitaine SARL", "city": "Bordeaux", "street": "44 Rue Sainte-Catherine, 33000"},
            {"name": "BureauModerne Nantes", "biz": "BureauModerne Ouest SAS", "city": "Nantes", "street": "10 Rue Crebillon, 44000"},
            {"name": "SportElite Strasbourg", "biz": "SportElite Est SARL", "city": "Strasbourg", "street": "7 Rue Grandes Arcades, 67000"},
            {"name": "CuisinePro Nice", "biz": "CuisinePro Cote Azur SARL", "city": "Nice", "street": "28 Av Jean Medecin, 06000"},
            {"name": "LumiereDesign Lille", "biz": "LumiereDesign Nord SAS", "city": "Lille", "street": "15 Rue Faidherbe, 59000"},
            {"name": "JardinTech Montpellier", "biz": "JardinTech Sud SARL", "city": "Montpellier", "street": "22 Rue de la Loge, 34000"},
        ],
    },
    "amazon.it": {
        "country": "Italy",
        "domain": "amazon.it",
        "vat_prefix": "IT",
        "phone_prefix": "+39",
        "email_tld": "it",
        "sellers": [
            {"name": "TecnoItalia Roma", "biz": "TecnoItalia S.r.l.", "city": "Roma", "street": "Via del Corso 125, 00186"},
            {"name": "ElettroShop Milano", "biz": "ElettroShop S.p.A.", "city": "Milano", "street": "Corso Buenos Aires 47, 20124"},
            {"name": "DigitalStore Napoli", "biz": "DigitalStore Campania S.r.l.", "city": "Napoli", "street": "Via Toledo 30, 80134"},
            {"name": "CasaModerna Torino", "biz": "CasaModerna Piemonte S.r.l.", "city": "Torino", "street": "Via Roma 88, 10121"},
            {"name": "AudioPerfetto Firenze", "biz": "AudioPerfetto Toscana S.r.l.", "city": "Firenze", "street": "Via dei Calzaiuoli 15, 50122"},
            {"name": "UfficioPlus Bologna", "biz": "UfficioPlus Emilia S.r.l.", "city": "Bologna", "street": "Via Indipendenza 42, 40121"},
            {"name": "SportInnovazione Venezia", "biz": "SportInnovazione Veneto S.r.l.", "city": "Venezia", "street": "Calle Larga XXII Marzo, 30124"},
            {"name": "CucinaTop Palermo", "biz": "CucinaTop Sicilia S.r.l.", "city": "Palermo", "street": "Via Maqueda 66, 90133"},
            {"name": "LuceDesign Genova", "biz": "LuceDesign Liguria S.r.l.", "city": "Genova", "street": "Via XX Settembre 33, 16121"},
            {"name": "GiardinoTech Bari", "biz": "GiardinoTech Puglia S.r.l.", "city": "Bari", "street": "Via Sparano 78, 70121"},
        ],
    },
}

# ═══════════════════════════════════════════════════
# HELPER FUNCTIONS
# ═══════════════════════════════════════════════════

def generate_asin():
    return "B0" + ''.join(random.choices(string.ascii_uppercase + string.digits, k=8))

def generate_seller_id():
    return ''.join(random.choices(string.ascii_uppercase + string.digits, k=13))

def generate_mock_seller(marketplace, seller_data=None, product_title=None, asin=None):
    config = MARKETPLACE_CONFIG[marketplace]
    if seller_data is None:
        seller_data = random.choice(config["sellers"])
    if product_title is None:
        product_title = random.choice(PRODUCT_TITLES)
    if asin is None:
        asin = generate_asin()

    words = product_title.split()
    short_title = ' '.join(words[:random.randint(5, min(8, len(words)))])
    seller_id = generate_seller_id()

    has_phone = random.random() > 0.3
    has_email = random.random() > 0.4
    has_vat = random.random() > 0.2
    has_reg = random.random() > 0.25

    phone = ""
    if has_phone:
        phone = f"{config['phone_prefix']} {random.randint(100, 999)} {random.randint(1000000, 9999999)}"

    email = ""
    if has_email:
        biz_key = seller_data["biz"].lower().split()[0]
        for ch in ".,()":
            biz_key = biz_key.replace(ch, "")
        email = f"info@{biz_key}.{config['email_tld']}"

    vat = ""
    if has_vat:
        vat = f"{config['vat_prefix']}{random.randint(100000000, 999999999)}"

    reg = ""
    if has_reg:
        if marketplace == "amazon.de":
            reg = f"HRB {random.randint(100000, 999999)}"
        elif marketplace == "amazon.es":
            reg = f"CIF {random.choice('ABCDEFGH')}{random.randint(10000000, 99999999)}"
        elif marketplace == "amazon.fr":
            reg = f"SIRET {random.randint(10000000000000, 99999999999999)}"
        elif marketplace == "amazon.it":
            reg = f"REA {random.choice(['MI', 'RM', 'NA', 'TO', 'FI', 'BO'])} {random.randint(100000, 999999)}"

    address = f"{seller_data['street']}, {seller_data['city']}, {config['country']}"
    now = datetime.now(timezone.utc).isoformat()

    seller = {
        "id": str(uuid.uuid4()),
        "marketplace": marketplace,
        "seller_name": seller_data["name"],
        "seller_profile_url": f"https://www.{config['domain']}/sp?seller={seller_id}",
        "product_title_short": short_title,
        "asin": asin,
        "product_url": f"https://www.{config['domain']}/dp/{asin}",
        "brand_if_visible": seller_data["biz"].split()[0] if random.random() > 0.3 else "",
        "seller_type_if_visible": random.choice(["Marketplace", "Third-Party", "FBA", "FBM"]),
        "business_name_if_visible": seller_data["biz"],
        "business_address_if_visible": address,
        "registration_number_if_visible": reg,
        "vat_number_if_visible": vat,
        "country_if_visible": config["country"],
        "public_phone_if_visible": phone,
        "public_email_if_visible": email,
        "collected_at": now,
        "first_seen_at": now,
        "last_seen_at": now,
        "source_page_type": random.choice(["seller_info", "product_page", "storefront"]),
        "notes": "",
        "status": "confirmed",
        "quality_score": 0,
        "flagged": False,
        "flagged_reason": "",
    }
    seller["quality_score"] = calculate_lead_score(seller)
    return seller

async def add_log(session_id, marketplace, action, detail, level="info"):
    entry = {
        "id": str(uuid.uuid4()),
        "timestamp": datetime.now(timezone.utc).isoformat(),
        "session_id": session_id,
        "marketplace": marketplace,
        "action": action,
        "detail": detail,
        "level": level,
    }
    await db.scan_log.insert_one(entry)

# ═══════════════════════════════════════════════════
# MOCK WORKER
# ═══════════════════════════════════════════════════

class MockWorker:
    def __init__(self):
        self._task = None
        self._running = False
        self.current_session_id = None
        self.current_marketplace = ""
        self.products_scanned = 0
        self.sellers_found = 0
        self.duplicates_skipped = 0
        self.started_at = None

    @property
    def is_running(self):
        return self._running

    def get_status(self):
        elapsed = 0
        if self.started_at and self._running:
            try:
                start = datetime.fromisoformat(self.started_at)
                elapsed = (datetime.now(timezone.utc) - start).total_seconds()
            except Exception:
                elapsed = 0
        sellers_per_hour = round(self.sellers_found / elapsed * 3600, 1) if elapsed > 0 else 0
        return {
            "is_running": self._running,
            "session_id": self.current_session_id,
            "marketplace": self.current_marketplace,
            "products_scanned": self.products_scanned,
            "sellers_found": self.sellers_found,
            "duplicates_skipped": self.duplicates_skipped,
            "started_at": self.started_at,
            "elapsed_seconds": round(elapsed),
            "sellers_per_hour": sellers_per_hour,
        }

    async def start(self, config: WorkerStartRequest):
        if self._running:
            raise HTTPException(400, "Worker already running")
        self._running = True
        self.products_scanned = 0
        self.sellers_found = 0
        self.duplicates_skipped = 0
        self.current_marketplace = config.marketplace
        self.started_at = datetime.now(timezone.utc).isoformat()

        session = {
            "id": str(uuid.uuid4()),
            "marketplace": config.marketplace,
            "mode": config.mode,
            "target_count": config.target_count,
            "min_delay": config.min_delay,
            "max_delay": config.max_delay,
            "manual_review": config.manual_review,
            "status": "running",
            "started_at": self.started_at,
            "stopped_at": None,
            "products_scanned": 0,
            "sellers_found": 0,
            "duplicates_skipped": 0,
            "last_asin": "",
        }
        self.current_session_id = session["id"]
        await db.sessions.insert_one(session)
        await add_log(session["id"], config.marketplace, "worker_start", f"Started scanning {config.marketplace}", "info")
        self._task = asyncio.create_task(self._run_loop(config))
        return {"session_id": session["id"], "status": "started"}

    async def stop(self):
        was_running = self._running
        self._running = False
        if self._task:
            self._task.cancel()
            try:
                await self._task
            except asyncio.CancelledError:
                pass
            self._task = None
        if self.current_session_id and was_running:
            await db.sessions.update_one(
                {"id": self.current_session_id},
                {"$set": {"status": "stopped", "stopped_at": datetime.now(timezone.utc).isoformat(),
                           "products_scanned": self.products_scanned, "sellers_found": self.sellers_found,
                           "duplicates_skipped": self.duplicates_skipped}}
            )
            await add_log(self.current_session_id, self.current_marketplace, "worker_stop",
                          f"Stopped. {self.sellers_found} sellers found, {self.duplicates_skipped} duplicates.", "info")

    async def _run_loop(self, config):
        session_id = self.current_session_id
        marketplace = config.marketplace
        used_names = set()
        try:
            while self._running:
                if config.mode == "target" and self.sellers_found >= config.target_count:
                    await db.sessions.update_one(
                        {"id": session_id},
                        {"$set": {"status": "completed", "stopped_at": datetime.now(timezone.utc).isoformat()}}
                    )
                    await add_log(session_id, marketplace, "target_reached",
                                  f"Target of {config.target_count} sellers reached!", "success")
                    self._running = False
                    break

                seller_data = random.choice(MARKETPLACE_CONFIG[marketplace]["sellers"])
                product_title = random.choice(PRODUCT_TITLES)
                asin = generate_asin()
                self.products_scanned += 1

                short = ' '.join(product_title.split()[:5])
                await add_log(session_id, marketplace, "scan_product", f"Scanning {asin}: {short}...", "info")

                delay = random.uniform(2, 4)
                await asyncio.sleep(delay)

                existing = await db.sellers.find_one(
                    {"seller_name": seller_data["name"], "marketplace": marketplace}, {"_id": 0}
                )
                is_dup = existing is not None or seller_data["name"] in used_names

                if is_dup:
                    self.duplicates_skipped += 1
                    await add_log(session_id, marketplace, "duplicate_skip",
                                  f"Duplicate: {seller_data['name']}", "warning")
                else:
                    seller = generate_mock_seller(marketplace, seller_data, product_title, asin)
                    if config.manual_review:
                        seller["status"] = "pending_review"
                        await add_log(session_id, marketplace, "pending_review",
                                      f"Pending review: {seller['seller_name']}", "info")
                    else:
                        await add_log(session_id, marketplace, "seller_saved",
                                      f"New seller: {seller['seller_name']} ({seller['country_if_visible']})", "success")
                    await db.sellers.insert_one(seller)
                    used_names.add(seller_data["name"])
                    self.sellers_found += 1

                await db.sessions.update_one(
                    {"id": session_id},
                    {"$set": {"products_scanned": self.products_scanned, "sellers_found": self.sellers_found,
                              "duplicates_skipped": self.duplicates_skipped, "last_asin": asin}}
                )
                await db.processed_urls.insert_one({
                    "url": f"https://www.{MARKETPLACE_CONFIG[marketplace]['domain']}/dp/{asin}",
                    "asin": asin, "session_id": session_id,
                    "processed_at": datetime.now(timezone.utc).isoformat(),
                })
        except asyncio.CancelledError:
            pass
        except Exception as e:
            logger.error(f"Worker error: {e}")
            await add_log(session_id, marketplace, "error", str(e), "error")
        finally:
            self._running = False

worker = MockWorker()

# ═══════════════════════════════════════════════════
# API ROUTES
# ═══════════════════════════════════════════════════

@api_router.get("/")
async def root():
    return {"message": "SellerRadar Pro API"}

@api_router.get("/stats")
async def get_stats():
    total_sellers = await db.sellers.count_documents({})
    confirmed = await db.sellers.count_documents({"status": "confirmed"})
    pending = await db.sellers.count_documents({"status": "pending_review"})
    flagged = await db.sellers.count_documents({"flagged": True})
    products_scanned = await db.processed_urls.count_documents({})
    status = worker.get_status()
    avg_score_agg = await db.sellers.aggregate([{"$group": {"_id": None, "avg": {"$avg": {"$ifNull": ["$quality_score", 0]}}}}]).to_list(1)
    avg_score = round(avg_score_agg[0]["avg"] or 0, 1) if avg_score_agg else 0
    return {
        "total_sellers": total_sellers,
        "confirmed_sellers": confirmed,
        "pending_review": pending,
        "flagged_count": flagged,
        "avg_quality_score": avg_score,
        "duplicates_skipped": status["duplicates_skipped"] if status["is_running"] else await db.sellers.count_documents({"status": "duplicate"}),
        "products_scanned": status["products_scanned"] if status["is_running"] else products_scanned,
        "sellers_per_hour": status["sellers_per_hour"],
        "is_running": status["is_running"],
        "run_status": "running" if status["is_running"] else "idle",
    }

@api_router.get("/sellers")
async def get_sellers(
    marketplace: Optional[str] = None,
    country: Optional[str] = None,
    has_phone: Optional[bool] = None,
    has_email: Optional[bool] = None,
    has_vat: Optional[bool] = None,
    has_registration: Optional[bool] = None,
    status: Optional[str] = None,
    search: Optional[str] = None,
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    min_score: Optional[int] = None,
    max_score: Optional[int] = None,
    flagged: Optional[bool] = None,
    page: int = 1,
    limit: int = 20,
    sort_by: str = "collected_at",
    sort_order: str = "desc",
):
    query = {}
    if marketplace:
        query["marketplace"] = marketplace
    if country:
        query["country_if_visible"] = {"$regex": country, "$options": "i"}
    if has_phone:
        query["public_phone_if_visible"] = {"$ne": ""}
    if has_email:
        query["public_email_if_visible"] = {"$ne": ""}
    if has_vat:
        query["vat_number_if_visible"] = {"$ne": ""}
    if has_registration:
        query["registration_number_if_visible"] = {"$ne": ""}
    if status:
        query["status"] = status
    if flagged is not None:
        query["flagged"] = flagged
    if min_score is not None:
        query.setdefault("quality_score", {})["$gte"] = min_score
    if max_score is not None:
        query.setdefault("quality_score", {})["$lte"] = max_score
    if date_from:
        query.setdefault("collected_at", {})["$gte"] = date_from
    if date_to:
        query.setdefault("collected_at", {})["$lte"] = date_to
    if search:
        query["$or"] = [
            {"seller_name": {"$regex": search, "$options": "i"}},
            {"business_name_if_visible": {"$regex": search, "$options": "i"}},
            {"product_title_short": {"$regex": search, "$options": "i"}},
            {"vat_number_if_visible": {"$regex": search, "$options": "i"}},
            {"registration_number_if_visible": {"$regex": search, "$options": "i"}},
        ]

    sort_dir = -1 if sort_order == "desc" else 1
    skip_n = (page - 1) * limit
    total = await db.sellers.count_documents(query)
    sellers = await db.sellers.find(query, {"_id": 0}).sort(sort_by, sort_dir).skip(skip_n).limit(limit).to_list(limit)
    return {"sellers": sellers, "total": total, "page": page, "limit": limit, "total_pages": max(1, (total + limit - 1) // limit)}

@api_router.delete("/sellers/{seller_id}")
async def delete_seller(seller_id: str):
    result = await db.sellers.delete_one({"id": seller_id})
    if result.deleted_count == 0:
        raise HTTPException(404, "Seller not found")
    return {"message": "Seller deleted"}

@api_router.get("/activity-log")
async def get_activity_log(session_id: Optional[str] = None, limit: int = 100):
    query = {}
    if session_id:
        query["session_id"] = session_id
    entries = await db.scan_log.find(query, {"_id": 0}).sort("timestamp", -1).limit(limit).to_list(limit)
    entries.reverse()
    return {"entries": entries, "total": await db.scan_log.count_documents(query)}

@api_router.get("/sessions/latest")
async def get_latest_session():
    session = await db.sessions.find_one(
        {"status": {"$in": ["stopped", "paused"]}}, {"_id": 0}, sort=[("started_at", -1)]
    )
    return {"session": session}

@api_router.post("/worker/start")
async def start_worker(req: WorkerStartRequest):
    return await worker.start(req)

@api_router.post("/worker/stop")
async def stop_worker():
    await worker.stop()
    return {"status": "stopped"}

@api_router.get("/worker/status")
async def get_worker_status():
    return worker.get_status()

@api_router.get("/review/pending")
async def get_pending_reviews():
    items = await db.sellers.find({"status": "pending_review"}, {"_id": 0}).sort("collected_at", 1).limit(10).to_list(10)
    return items

@api_router.post("/review/{seller_id}/confirm")
async def confirm_review(seller_id: str):
    result = await db.sellers.update_one({"id": seller_id}, {"$set": {"status": "confirmed"}})
    if result.matched_count == 0:
        raise HTTPException(404, "Seller not found")
    return {"message": "Seller confirmed"}

@api_router.post("/review/{seller_id}/skip")
async def skip_review(seller_id: str):
    result = await db.sellers.update_one({"id": seller_id}, {"$set": {"status": "skipped"}})
    if result.matched_count == 0:
        raise HTTPException(404, "Seller not found")
    return {"message": "Seller skipped"}

@api_router.get("/settings")
async def get_settings():
    settings = await db.settings.find_one({"id": "app_settings"}, {"_id": 0})
    if not settings:
        default = {
            "id": "app_settings", "google_sheets_enabled": False, "google_sheets_url": "",
            "manual_review_enabled": False, "default_marketplace": "amazon.de",
            "default_min_delay": 25, "default_max_delay": 45,
            "default_mode": "continuous", "default_target_count": 50,
            "scoring_weights": DEFAULT_SCORING_WEIGHTS,
            "duplicate_rules": {"by_profile_url": True, "by_name_marketplace": True, "by_email": True, "by_phone": True, "by_vat_or_reg": True},
            "export_options": {"include_flagged_sheet": True, "include_analytics_sheet": True, "alternating_rows": True},
        }
        await db.settings.insert_one(default)
        return default
    return settings

@api_router.put("/settings")
async def update_settings(req: SettingsUpdate):
    update_data = {k: v for k, v in req.model_dump().items() if v is not None}
    if not update_data:
        raise HTTPException(400, "No fields to update")
    existing = await db.settings.find_one({"id": "app_settings"})
    if not existing:
        await db.settings.insert_one({"id": "app_settings", **update_data})
    else:
        await db.settings.update_one({"id": "app_settings"}, {"$set": update_data})
    return await db.settings.find_one({"id": "app_settings"}, {"_id": 0})

# ═══════════════ NEW ENDPOINTS ═══════════════

@api_router.get("/sellers/{seller_id}")
async def get_seller_detail(seller_id: str):
    seller = await db.sellers.find_one({"id": seller_id}, {"_id": 0})
    if not seller:
        raise HTTPException(404, "Seller not found")
    settings = await db.settings.find_one({"id": "app_settings"}, {"_id": 0})
    weights = settings.get("scoring_weights", DEFAULT_SCORING_WEIGHTS) if settings else DEFAULT_SCORING_WEIGHTS
    seller["score_breakdown"] = get_score_breakdown(seller, weights)
    return seller

@api_router.put("/sellers/{seller_id}/notes")
async def update_seller_notes(seller_id: str, req: NotesUpdate):
    result = await db.sellers.update_one({"id": seller_id}, {"$set": {"notes": req.notes}})
    if result.matched_count == 0:
        raise HTTPException(404, "Seller not found")
    return {"message": "Notes updated"}

@api_router.post("/sellers/{seller_id}/flag")
async def flag_seller(seller_id: str, req: FlagUpdate):
    result = await db.sellers.update_one({"id": seller_id}, {"$set": {"flagged": req.flagged, "flagged_reason": req.flagged_reason}})
    if result.matched_count == 0:
        raise HTTPException(404, "Seller not found")
    return {"message": "Seller flagged" if req.flagged else "Flag removed"}

@api_router.get("/analytics")
async def get_analytics():
    total = await db.sellers.count_documents({})
    if total == 0:
        return {"marketplace_distribution": [], "records_over_time": [], "data_quality": {}, "score_distribution": [], "duplicate_rate": 0, "avg_score": 0, "flagged_count": 0}

    mp_dist = await db.sellers.aggregate([
        {"$group": {"_id": "$marketplace", "count": {"$sum": 1}}}
    ]).to_list(100)

    records_time = await db.sellers.aggregate([
        {"$project": {"day": {"$substr": ["$collected_at", 0, 10]}}},
        {"$group": {"_id": "$day", "count": {"$sum": 1}}},
        {"$sort": {"_id": 1}}
    ]).to_list(100)

    has_email = await db.sellers.count_documents({"public_email_if_visible": {"$ne": ""}})
    has_phone = await db.sellers.count_documents({"public_phone_if_visible": {"$ne": ""}})
    has_vat = await db.sellers.count_documents({"vat_number_if_visible": {"$ne": ""}})
    has_reg = await db.sellers.count_documents({"registration_number_if_visible": {"$ne": ""}})
    has_addr = await db.sellers.count_documents({"business_address_if_visible": {"$ne": ""}})
    flagged = await db.sellers.count_documents({"flagged": True})

    score_agg = await db.sellers.aggregate([
        {"$bucket": {"groupBy": "$quality_score", "boundaries": [0, 20, 40, 60, 80, 101],
                      "default": "other", "output": {"count": {"$sum": 1}}}},
    ]).to_list(100)
    score_dist = []
    labels = ["0-19", "20-39", "40-59", "60-79", "80-100"]
    for i, b in enumerate(score_agg):
        score_dist.append({"range": labels[i] if i < len(labels) else "other", "count": b["count"]})

    avg_score_agg = await db.sellers.aggregate([
        {"$group": {"_id": None, "avg": {"$avg": "$quality_score"}}}
    ]).to_list(1)
    avg_score = round(avg_score_agg[0]["avg"], 1) if avg_score_agg else 0

    dup_pipeline = [
        {"$group": {"_id": {"name": "$seller_name", "mp": "$marketplace"}, "count": {"$sum": 1}}},
        {"$match": {"count": {"$gt": 1}}},
        {"$count": "total"}
    ]
    dup_result = await db.sellers.aggregate(dup_pipeline).to_list(1)
    dup_rate = round((dup_result[0]["total"] / total * 100) if dup_result else 0, 1)

    return {
        "marketplace_distribution": [{"name": d["_id"], "value": d["count"]} for d in mp_dist],
        "records_over_time": [{"date": d["_id"], "count": d["count"]} for d in records_time],
        "data_quality": {"total": total, "has_email": has_email, "has_phone": has_phone, "has_vat": has_vat, "has_registration": has_reg, "has_address": has_addr},
        "score_distribution": score_dist,
        "duplicate_rate": dup_rate,
        "avg_score": avg_score,
        "flagged_count": flagged,
    }

@api_router.get("/saved-views")
async def get_saved_views():
    views = await db.saved_views.find({}, {"_id": 0}).to_list(100)
    return views

@api_router.post("/saved-views")
async def create_saved_view(req: SavedViewCreate):
    view = {"id": str(uuid.uuid4()), "name": req.name, "filters": req.filters, "created_at": datetime.now(timezone.utc).isoformat()}
    await db.saved_views.insert_one(view)
    return view

@api_router.delete("/saved-views/{view_id}")
async def delete_saved_view(view_id: str):
    await db.saved_views.delete_one({"id": view_id})
    return {"message": "View deleted"}

@api_router.get("/review/queue")
async def get_review_queue(page: int = 1, limit: int = 1):
    total = await db.sellers.count_documents({"status": "pending_review"})
    skip_n = (page - 1) * limit
    items = await db.sellers.find({"status": "pending_review"}, {"_id": 0}).sort("collected_at", 1).skip(skip_n).limit(limit).to_list(limit)
    for item in items:
        item["score_breakdown"] = get_score_breakdown(item)
    return {"items": items, "total": total, "page": page}

@api_router.post("/review/{seller_id}/flag-duplicate")
async def flag_review_duplicate(seller_id: str):
    result = await db.sellers.update_one({"id": seller_id}, {"$set": {"status": "duplicate", "flagged": True, "flagged_reason": "Flagged as duplicate during review"}})
    if result.matched_count == 0:
        raise HTTPException(404, "Seller not found")
    return {"message": "Flagged as duplicate"}

# ═══════════════════════════════════════════════════
# EXPORTS
# ═══════════════════════════════════════════════════

@api_router.get("/export/excel")
async def export_excel():
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="1a2230", end_color="1a2230", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center")

    # Sheet 1: Sellers
    ws1 = wb.active
    ws1.title = "Sellers"
    cols = ["Marketplace", "Seller Name", "Business Name", "Country", "Quality Score", "VAT Number", "Registration",
            "Phone", "Email", "Product Title", "ASIN", "Brand", "Seller Type", "Address",
            "Profile URL", "Collected At", "Status", "Flagged"]
    for c, h in enumerate(cols, 1):
        cell = ws1.cell(row=1, column=c, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    sellers = await db.sellers.find({}, {"_id": 0}).sort("collected_at", -1).to_list(50000)
    alt_fill = PatternFill(start_color="0f1419", end_color="0f1419", fill_type="solid")
    for r, s in enumerate(sellers, 2):
        ws1.cell(row=r, column=1, value=s.get("marketplace", ""))
        ws1.cell(row=r, column=2, value=s.get("seller_name", ""))
        ws1.cell(row=r, column=3, value=s.get("business_name_if_visible", ""))
        ws1.cell(row=r, column=4, value=s.get("country_if_visible", ""))
        ws1.cell(row=r, column=5, value=s.get("quality_score", 0))
        ws1.cell(row=r, column=6, value=s.get("vat_number_if_visible", ""))
        ws1.cell(row=r, column=7, value=s.get("registration_number_if_visible", ""))
        ws1.cell(row=r, column=8, value=s.get("public_phone_if_visible", ""))
        ws1.cell(row=r, column=9, value=s.get("public_email_if_visible", ""))
        ws1.cell(row=r, column=10, value=s.get("product_title_short", ""))
        ws1.cell(row=r, column=11, value=s.get("asin", ""))
        ws1.cell(row=r, column=12, value=s.get("brand_if_visible", ""))
        ws1.cell(row=r, column=13, value=s.get("seller_type_if_visible", ""))
        ws1.cell(row=r, column=14, value=s.get("business_address_if_visible", ""))
        ws1.cell(row=r, column=15, value=s.get("seller_profile_url", ""))
        ws1.cell(row=r, column=16, value=s.get("collected_at", "")[:19].replace("T", " ") if s.get("collected_at") else "")
        ws1.cell(row=r, column=17, value=s.get("status", ""))
        ws1.cell(row=r, column=18, value="Yes" if s.get("flagged") else "")
        if r % 2 == 0:
            for c in range(1, len(cols) + 1):
                ws1.cell(row=r, column=c).fill = alt_fill

    ws1.auto_filter.ref = ws1.dimensions
    ws1.freeze_panes = "A2"
    for c in range(1, len(cols) + 1):
        max_len = max((len(str(ws1.cell(row=r, column=c).value or "")) for r in range(1, min(ws1.max_row + 1, 100))), default=10)
        ws1.column_dimensions[get_column_letter(c)].width = min(max_len + 3, 50)

    # Sheet 2: Scan Log
    ws2 = wb.create_sheet("Scan Log")
    log_cols = ["Timestamp", "Marketplace", "Action", "Detail", "Level", "Session ID"]
    for c, h in enumerate(log_cols, 1):
        cell = ws2.cell(row=1, column=c, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    logs = await db.scan_log.find({}, {"_id": 0}).sort("timestamp", -1).limit(5000).to_list(5000)
    for r, log_entry in enumerate(logs, 2):
        ws2.cell(row=r, column=1, value=log_entry.get("timestamp", "")[:19].replace("T", " ") if log_entry.get("timestamp") else "")
        ws2.cell(row=r, column=2, value=log_entry.get("marketplace", ""))
        ws2.cell(row=r, column=3, value=log_entry.get("action", ""))
        ws2.cell(row=r, column=4, value=log_entry.get("detail", ""))
        ws2.cell(row=r, column=5, value=log_entry.get("level", ""))
        ws2.cell(row=r, column=6, value=log_entry.get("session_id", ""))

    ws2.auto_filter.ref = ws2.dimensions
    ws2.freeze_panes = "A2"
    for c in range(1, len(log_cols) + 1):
        max_len = max((len(str(ws2.cell(row=r, column=c).value or "")) for r in range(1, min(ws2.max_row + 1, 100))), default=10)
        ws2.column_dimensions[get_column_letter(c)].width = min(max_len + 3, 50)

    # Sheet 3: Duplicate Summary
    ws3 = wb.create_sheet("Duplicate Summary")
    dup_cols = ["Seller Name", "Occurrences", "Marketplaces"]
    for c, h in enumerate(dup_cols, 1):
        cell = ws3.cell(row=1, column=c, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    pipeline = [
        {"$group": {"_id": "$seller_name", "count": {"$sum": 1}, "marketplaces": {"$addToSet": "$marketplace"}}},
        {"$match": {"count": {"$gt": 1}}},
        {"$sort": {"count": -1}}
    ]
    dupes = await db.sellers.aggregate(pipeline).to_list(1000)
    for r, d in enumerate(dupes, 2):
        ws3.cell(row=r, column=1, value=d.get("_id", ""))
        ws3.cell(row=r, column=2, value=d.get("count", 0))
        ws3.cell(row=r, column=3, value=", ".join(d.get("marketplaces", [])))

    ws3.auto_filter.ref = ws3.dimensions
    ws3.freeze_panes = "A2"
    for c in range(1, len(dup_cols) + 1):
        max_len = max((len(str(ws3.cell(row=r, column=c).value or "")) for r in range(1, min(ws3.max_row + 1, 100))), default=10)
        ws3.column_dimensions[get_column_letter(c)].width = min(max_len + 3, 40)

    # Sheet 4: Flagged Records
    ws4 = wb.create_sheet("Flagged Records")
    flag_cols = ["Seller Name", "Marketplace", "Country", "Score", "Reason", "Status"]
    for c, h in enumerate(flag_cols, 1):
        cell = ws4.cell(row=1, column=c, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
    flagged_sellers = [s for s in sellers if s.get("flagged")]
    for r, s in enumerate(flagged_sellers, 2):
        ws4.cell(row=r, column=1, value=s.get("seller_name", ""))
        ws4.cell(row=r, column=2, value=s.get("marketplace", ""))
        ws4.cell(row=r, column=3, value=s.get("country_if_visible", ""))
        ws4.cell(row=r, column=4, value=s.get("quality_score", 0))
        ws4.cell(row=r, column=5, value=s.get("flagged_reason", ""))
        ws4.cell(row=r, column=6, value=s.get("status", ""))
    ws4.auto_filter.ref = ws4.dimensions
    ws4.freeze_panes = "A2"
    for c in range(1, len(flag_cols) + 1):
        max_len = max((len(str(ws4.cell(row=r, column=c).value or "")) for r in range(1, min(ws4.max_row + 1, 50))), default=10)
        ws4.column_dimensions[get_column_letter(c)].width = min(max_len + 3, 40)

    # Sheet 5: Analytics Summary
    ws5 = wb.create_sheet("Analytics Summary")
    title_font = Font(bold=True, color="3b82f6", size=14)
    ws5.cell(row=1, column=1, value="SellerRadar Pro - Analytics Summary").font = title_font
    ws5.cell(row=2, column=1, value=f"Generated: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}")
    ws5.cell(row=4, column=1, value="Metric").font = header_font
    ws5.cell(row=4, column=2, value="Value").font = header_font
    ws5.cell(row=4, column=1).fill = header_fill
    ws5.cell(row=4, column=2).fill = header_fill
    metrics = [
        ("Total Sellers", len(sellers)),
        ("Avg Quality Score", round(sum(s.get("quality_score", 0) for s in sellers) / max(len(sellers), 1), 1)),
        ("Has Email", sum(1 for s in sellers if s.get("public_email_if_visible"))),
        ("Has Phone", sum(1 for s in sellers if s.get("public_phone_if_visible"))),
        ("Has VAT", sum(1 for s in sellers if s.get("vat_number_if_visible"))),
        ("Flagged Records", len(flagged_sellers)),
    ]
    for i, (m, v) in enumerate(metrics, 5):
        ws5.cell(row=i, column=1, value=m)
        ws5.cell(row=i, column=2, value=v)
    ws5.column_dimensions["A"].width = 25
    ws5.column_dimensions["B"].width = 15

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": "attachment; filename=sellerradar_export.xlsx"})

@api_router.get("/export/csv")
async def export_csv():
    sellers = await db.sellers.find({}, {"_id": 0}).sort("collected_at", -1).to_list(50000)
    output = io.StringIO()
    fields = ["marketplace", "seller_name", "business_name_if_visible", "country_if_visible",
              "quality_score", "vat_number_if_visible", "registration_number_if_visible", "public_phone_if_visible",
              "public_email_if_visible", "product_title_short", "asin", "brand_if_visible",
              "seller_type_if_visible", "business_address_if_visible", "seller_profile_url",
              "collected_at", "status", "flagged"]
    w = csv_module.DictWriter(output, fieldnames=fields, extrasaction='ignore')
    w.writeheader()
    for s in sellers:
        w.writerow(s)
    output.seek(0)
    return StreamingResponse(io.BytesIO(output.getvalue().encode('utf-8')), media_type="text/csv",
                             headers={"Content-Disposition": "attachment; filename=sellerradar_export.csv"})

# ═══════════════════════════════════════════════════
# SEED & CLEAR
# ═══════════════════════════════════════════════════

@api_router.post("/seed")
async def seed_data():
    await db.sellers.delete_many({})
    await db.scan_log.delete_many({})
    await db.sessions.delete_many({})
    await db.processed_urls.delete_many({})
    await db.settings.delete_many({})
    await db.saved_views.delete_many({})

    now = datetime.now(timezone.utc)
    sellers_created = 0
    session_id = str(uuid.uuid4())

    completeness_profiles = [
        {"phone": True, "email": True, "vat": True, "reg": True, "addr": True},
        {"phone": True, "email": True, "vat": True, "reg": False, "addr": True},
        {"phone": True, "email": False, "vat": True, "reg": True, "addr": True},
        {"phone": False, "email": True, "vat": True, "reg": False, "addr": True},
        {"phone": True, "email": True, "vat": False, "reg": False, "addr": True},
        {"phone": False, "email": False, "vat": True, "reg": True, "addr": True},
        {"phone": True, "email": False, "vat": False, "reg": False, "addr": True},
        {"phone": False, "email": True, "vat": False, "reg": False, "addr": False},
        {"phone": False, "email": False, "vat": False, "reg": False, "addr": True},
        {"phone": False, "email": False, "vat": False, "reg": False, "addr": False},
    ]

    for marketplace, config in MARKETPLACE_CONFIG.items():
        all_sellers = config["sellers"]
        for idx, seller_data in enumerate(all_sellers):
            num_records = random.choice([3, 3, 4, 4, 5])
            for j in range(num_records):
                seller = generate_mock_seller(marketplace, seller_data)
                profile = random.choice(completeness_profiles)
                if not profile["phone"]: seller["public_phone_if_visible"] = ""
                if not profile["email"]: seller["public_email_if_visible"] = ""
                if not profile["vat"]: seller["vat_number_if_visible"] = ""
                if not profile["reg"]: seller["registration_number_if_visible"] = ""
                if not profile["addr"]: seller["business_address_if_visible"] = ""
                seller["quality_score"] = calculate_lead_score(seller)
                days_ago = random.randint(0, 45)
                collected = (now - timedelta(days=days_ago, hours=random.randint(0, 23))).isoformat()
                seller["collected_at"] = collected
                seller["first_seen_at"] = collected
                seller["last_seen_at"] = collected
                if random.random() < 0.06:
                    seller["flagged"] = True
                    seller["flagged_reason"] = random.choice(["Suspicious data pattern", "Possible duplicate entry", "Needs manual verification", "Incomplete business info"])
                if random.random() < 0.08:
                    seller["status"] = "pending_review"
                elif random.random() < 0.03:
                    seller["status"] = "skipped"
                await db.sellers.insert_one(seller)
                sellers_created += 1

    actions_pool = ["scan_product", "seller_saved", "duplicate_skip", "scan_product", "seller_saved", "scan_product", "pending_review"]
    for i in range(120):
        mp = random.choice(list(MARKETPLACE_CONFIG.keys()))
        action = random.choice(actions_pool)
        level = "success" if action == "seller_saved" else "warning" if action == "duplicate_skip" else "info"
        ts = (now - timedelta(hours=random.randint(0, 72), minutes=random.randint(0, 59))).isoformat()
        detail = f"{'New seller found' if action == 'seller_saved' else 'Scanning product' if action == 'scan_product' else 'Duplicate detected' if action == 'duplicate_skip' else 'Pending review'} ({generate_asin()})"
        await db.scan_log.insert_one({
            "id": str(uuid.uuid4()), "timestamp": ts, "session_id": session_id,
            "marketplace": mp, "action": action, "detail": detail, "level": level,
        })

    session = {
        "id": session_id, "marketplace": "amazon.de", "mode": "continuous",
        "target_count": 0, "min_delay": 25, "max_delay": 45, "manual_review": False,
        "status": "stopped", "started_at": (now - timedelta(hours=3)).isoformat(),
        "stopped_at": (now - timedelta(hours=1)).isoformat(),
        "products_scanned": 340, "sellers_found": sellers_created, "duplicates_skipped": 45, "last_asin": generate_asin(),
    }
    await db.sessions.insert_one(session)

    for i in range(340):
        asin = generate_asin()
        mp = random.choice(list(MARKETPLACE_CONFIG.keys()))
        await db.processed_urls.insert_one({
            "url": f"https://www.{MARKETPLACE_CONFIG[mp]['domain']}/dp/{asin}",
            "asin": asin, "session_id": session_id,
            "processed_at": (now - timedelta(hours=random.randint(0, 72), minutes=random.randint(0, 59))).isoformat(),
        })

    await db.settings.insert_one({
        "id": "app_settings", "google_sheets_enabled": False, "google_sheets_url": "",
        "manual_review_enabled": False, "default_marketplace": "amazon.de",
        "default_min_delay": 25, "default_max_delay": 45,
        "default_mode": "continuous", "default_target_count": 50,
        "scoring_weights": DEFAULT_SCORING_WEIGHTS,
        "duplicate_rules": {"by_profile_url": True, "by_name_marketplace": True, "by_email": True, "by_phone": True, "by_vat_or_reg": True},
        "export_options": {"include_flagged_sheet": True, "include_analytics_sheet": True, "alternating_rows": True},
    })

    quick_views = [
        {"id": str(uuid.uuid4()), "name": "Has Email", "filters": {"has_email": True}, "created_at": now.isoformat()},
        {"id": str(uuid.uuid4()), "name": "Has Phone", "filters": {"has_phone": True}, "created_at": now.isoformat()},
        {"id": str(uuid.uuid4()), "name": "High Quality", "filters": {"min_score": 80}, "created_at": now.isoformat()},
        {"id": str(uuid.uuid4()), "name": "Missing VAT", "filters": {"has_vat": False}, "created_at": now.isoformat()},
        {"id": str(uuid.uuid4()), "name": "Germany Only", "filters": {"marketplace": "amazon.de"}, "created_at": now.isoformat()},
        {"id": str(uuid.uuid4()), "name": "Flagged", "filters": {"flagged": True}, "created_at": now.isoformat()},
    ]
    await db.saved_views.insert_many(quick_views)

    return {"message": f"Demo data seeded: {sellers_created} sellers across 4 marketplaces", "sellers_count": sellers_created}

@api_router.post("/data/clear")
async def clear_data():
    await db.sellers.delete_many({})
    await db.scan_log.delete_many({})
    await db.sessions.delete_many({})
    await db.processed_urls.delete_many({})
    return {"message": "All data cleared"}

# ═══════════════════════════════════════════════════
# APP SETUP
# ═══════════════════════════════════════════════════

app.include_router(api_router)

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=os.environ.get('CORS_ORIGINS', '*').split(','),
    allow_methods=["*"],
    allow_headers=["*"],
)

@app.on_event("shutdown")
async def shutdown_db_client():
    await worker.stop()
    client.close()
